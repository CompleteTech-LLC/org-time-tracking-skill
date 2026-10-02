#!/usr/bin/env python3
"""Render a self-contained, branded light/dark time-tracking HTML report from a RECALCULATED org-time-tracking workbook.

Usage:
  python build_report.py --config config.json [--workbook book.xlsx] [--out report.html] [--theme auto|light|dark]

The workbook must carry cached values (recalculate it in Excel or LibreOffice first: scripts/recalc_excel.ps1). The
report is a single HTML file: no external fonts, scripts or images; the logo (local file only) is inlined as a data URI;
a Content-Security-Policy meta tag forbids everything else. Branding is neutral unless a preset or explicit values are set
(see scripts/branding.py and templates/branding.*.json).
"""
from __future__ import annotations

import argparse
import base64
import calendar
import datetime as dt
import html
import json
import mimetypes
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import branding as brandlib  # noqa: E402


def esc(value) -> str:
    return html.escape("" if value is None else str(value))


def hm(minutes) -> str:
    minutes = int(round(minutes or 0))
    return f"{minutes // 60}:{minutes % 60:02d}"


def table(ws) -> list[dict]:
    """Rows of a tab whose first row is a header; stops at the first row with an empty first cell."""
    headers = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] in (None, ""):
            break
        rows.append({h: v for h, v in zip(headers, r) if h})
    return rows


def css(brand: dict, accent_on: str) -> str:
    def block(tokens: dict) -> str:
        return "".join(f"--{k}:{v};" for k, v in tokens.items())
    font = brand.get("font") or "Arial"
    text_light = brandlib.readable_on(brand["accent"], brand["light"]["surface"])
    text_dark = brandlib.readable_on(brand["accent"], brand["dark"]["surface"])
    return (
        f":root{{{block(brand['light'])}--accent:{brand['accent']};--on-accent:{accent_on};--accent-text:{text_light};}}"
        f"@media (prefers-color-scheme:dark){{:root:not([data-theme=\"light\"]){{{block(brand['dark'])}--accent-text:{text_dark};}}}}"
        f":root[data-theme=\"dark\"]{{{block(brand['dark'])}--accent-text:{text_dark};}}"
        f"*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 '{esc(font)}',system-ui,sans-serif}}"
        "main{max-width:1080px;margin:0 auto;padding:24px 16px 48px}"
        "header.brand{display:flex;gap:16px;align-items:center;padding:20px 0;border-bottom:3px solid var(--accent);margin-bottom:24px}"
        "header.brand img{height:56px;width:auto}.eyebrow{font-size:12px;letter-spacing:.12em;font-weight:700;color:var(--accent-text)}"
        "h1{margin:2px 0;font-size:26px}.tagline{color:var(--ink-3);margin:0}h2{margin:32px 0 10px;font-size:18px;color:var(--ink-2)}"
        ".tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}"
        ".tile{background:var(--surface);border:1px solid var(--line);border-radius:10px;padding:14px}"
        ".tile b{display:block;font-size:26px;color:var(--accent-text)}.tile span{color:var(--ink-3);font-size:13px}"
        "table{width:100%;border-collapse:collapse;background:var(--surface);border:1px solid var(--line);border-radius:8px;overflow:hidden}"
        "th{background:var(--accent);color:var(--on-accent);text-align:left;padding:8px 10px;font-size:13px}"
        "td{padding:8px 10px;border-top:1px solid var(--line);vertical-align:top;font-size:14px}tr:nth-child(even) td{background:var(--surface-2)}"
        "td.n,th.n{text-align:right;white-space:nowrap}.cal{display:grid;grid-template-columns:repeat(7,1fr);gap:4px;margin-bottom:8px}"
        ".cal div{border:1px solid var(--line);border-radius:6px;min-height:46px;padding:4px 6px;background:var(--surface);font-size:12px}"
        ".cal .dow{background:var(--surface-2);min-height:0;text-align:center;font-weight:700;color:var(--ink-2)}"
        ".cal .d{color:var(--ink-3);display:block;font-size:11px}.cal .m{font-weight:700;font-size:13px}"
        ".cal .hot{color:var(--on-accent)}.cal .hot .d{color:inherit;opacity:.8}.cal .blank{border:none;background:none}"
        ".note{color:var(--ink-3);font-size:13px}footer{margin-top:40px;padding-top:12px;border-top:1px solid var(--line);color:var(--ink-3);font-size:12px}"
        "@media print{body{background:#fff;color:#000}.tile,table,.cal div{break-inside:avoid}}"
    )


def build(cfg: dict, workbook: str, theme: str, config_dir: str) -> str:
    brand = brandlib.resolve(cfg.get("branding"), config_dir)
    accent_on = brandlib.on_colour(brand["accent"])
    scale_max = (cfg.get("style", {}).get("calendar_color_scale") or {}).get("max_min", 240)
    wb = load_workbook(workbook, data_only=True)
    for name in ("Summary", "Time Spent", "Daily Time", "Meetings", "Tasks"):
        if name not in wb.sheetnames:
            raise SystemExit(f"workbook has no '{name}' tab: build it with scripts/make_workbook.py")
    ts = wb["Time Spent"]
    if ts["D7"].value is None:
        raise SystemExit("workbook has no cached values: recalculate it first (scripts/recalc_excel.ps1 or LibreOffice)")

    # Summary metrics: label/value pairs under the banner header row "Metric"
    metrics, ws0 = [], wb["Summary"]
    start = next((r for r in range(1, 20) if ws0.cell(row=r, column=1).value == "Metric"), 5) + 1
    for r in range(start, start + 12):
        label, value = ws0.cell(row=r, column=1).value, ws0.cell(row=r, column=2).value
        if not label:
            break
        metrics.append((str(label).strip(), value))

    lines = [(ts.cell(row=r, column=1).value, ts.cell(row=r, column=2).value, ts.cell(row=r, column=4).value) for r in range(2, 10)]
    daily = {row["Date"].date(): row["Total (min)"] or 0 for row in table(wb["Daily Time"]) if isinstance(row.get("Date"), dt.datetime)}
    meetings = table(wb["Meetings"])
    tasks = [t for t in table(wb["Tasks"]) if isinstance(t.get("#"), (int, float))]
    org = cfg.get("org", {}).get("name", "")
    period = cfg.get("period", {})
    cfg_theme = theme or cfg.get("style", {}).get("theme", "auto")
    theme_attr = f' data-theme="{esc(cfg_theme)}"' if cfg_theme in ("light", "dark") else ""

    logo = ""
    if brand.get("logo_path"):
        mime = mimetypes.guess_type(brand["logo_path"])[0] or "image/png"
        data = base64.b64encode(open(brand["logo_path"], "rb").read()).decode("ascii")
        logo = f'<img src="data:{mime};base64,{data}" alt="{esc(brand.get("name") or "logo")} logo">'

    out = [f'<!doctype html><html lang="en"{theme_attr}><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width,initial-scale=1">',
           "<meta http-equiv=\"Content-Security-Policy\" content=\"default-src 'none'; style-src 'unsafe-inline'; img-src data:\">",
           f"<title>{esc(org)} time tracking report</title><style>{css(brand, accent_on)}</style></head><body><main>"]
    out.append('<header class="brand">' + logo + "<div>" + (f'<div class="eyebrow">{esc(brand["eyebrow"])}</div>' if brand.get("eyebrow") else "")
               + f"<h1>{esc(org)} time tracking</h1><p class=\"tagline\">{esc(brand.get('tagline') or '')}"
               + f" {esc(period.get('start', ''))} to {esc(period.get('end', ''))}</p></div></header>")

    tiles = [m for m in metrics if m[1] is not None][:10]
    out.append('<section class="tiles">' + "".join(
        f'<div class="tile"><b>{esc(f"{v:.1f}" if isinstance(v, float) else v)}</b><span>{esc(k)}</span></div>' for k, v in tiles) + "</section>")

    out.append("<h2>Billed time</h2><table><tr><th>Line item</th><th class=\"n\">Count</th><th class=\"n\">Minutes</th><th class=\"n\">h:mm</th></tr>")
    for label, count, minutes in lines:
        if label:
            out.append(f'<tr><td>{esc(label)}</td><td class="n">{esc(count)}</td><td class="n">{esc(round(minutes or 0))}</td><td class="n">{hm(minutes)}</td></tr>')
    out.append("</table>")

    out.append("<h2>Time per day</h2>")
    months = sorted({(d.year, d.month) for d in daily})
    for y, m in months:
        total = sum(v for d, v in daily.items() if (d.year, d.month) == (y, m))
        out.append(f"<h3>{calendar.month_name[m]} {y} · {hm(total)} h:mm</h3><div class=\"cal\">")
        out.extend(f'<div class="dow">{d}</div>' for d in ("Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"))
        for week in calendar.Calendar(firstweekday=6).monthdatescalendar(y, m):
            for day in week:
                if day.month != m:
                    out.append('<div class="blank"></div>')
                    continue
                v = daily.get(day, 0)
                pct = min(100, round(100 * v / scale_max)) if scale_max else 0
                style = f' style="background:color-mix(in srgb,var(--accent) {pct}%,var(--surface))"' if v else ""
                cls = " hot" if pct >= 55 else ""
                out.append(f'<div class="c{cls}"{style}><span class="d">{day.day}</span><span class="m">{esc(str(round(v)) + " m") if v else ""}</span></div>')
        out.append("</div>")
    out.append('<p class="note">Shade = billed minutes that day relative to the configured maximum; see the workbook (Daily Time, Day Detail) for the events.</p>')

    held = [r for r in meetings if str(r.get("Status", "")).startswith("Held")]
    out.append(f"<h2>Meetings counted as held ({len(held)})</h2><table><tr><th>Date</th><th>Meeting</th><th>Status</th><th class=\"n\">Teams card (min)</th><th>Evidence</th></tr>")
    for r in held:
        out.append(f"<tr><td>{esc(r['Date'].date() if isinstance(r.get('Date'), dt.datetime) else r.get('Date'))}</td><td>{esc(r.get('Meeting'))}</td>"
                   f"<td>{esc(r.get('Status'))}</td><td class=\"n\">{esc(r.get('Teams actual (min)'))}</td><td>{esc(r.get('Evidence'))}</td></tr>")
    out.append("</table>")
    labels = cfg.get("meetings", {}).get("labels", {})
    not_counted = tuple(labels.get(k, d) for k, d in (("not_held", "Not held (no conversation that day)"), ("likely_not_held", "Likely not held (chat)"), ("cancelled", "Cancelled")))
    skipped = [r for r in meetings if str(r.get("Status", "")) in not_counted]
    out.append(f"<h2>Scheduled but not counted ({len(skipped)})</h2><table><tr><th>Date</th><th>Meeting</th><th>Status</th></tr>")
    for r in skipped:
        out.append(f"<tr><td>{esc(r['Date'].date() if isinstance(r.get('Date'), dt.datetime) else r.get('Date'))}</td><td>{esc(r.get('Meeting'))}</td><td>{esc(r.get('Status'))}</td></tr>")
    out.append('</table><p class="note">With the default held-meeting rule, a scheduled call with no Teams conversation that day is not counted, even if the calendar shows it.</p>')

    out.append("<h2>Tasks and effort estimates</h2><table><tr><th>Task</th><th>Status</th><th class=\"n\">Low h</th><th class=\"n\">High h</th><th class=\"n\">Likely h</th></tr>")
    for t in tasks:
        out.append(f"<tr><td>{esc(t.get('Task'))}</td><td>{esc(t.get('Status'))}</td><td class=\"n\">{esc(t.get('Est. low (h)'))}</td>"
                   f"<td class=\"n\">{esc(t.get('Est. high (h)'))}</td><td class=\"n\">{esc(round(t.get('Est. likely (h)') or 0, 1))}</td></tr>")
    out.append('</table><p class="note">Estimates from the scope described in the messages, not tracked time.</p>')

    footer = brand.get("footer") or ""
    contact = brand.get("contact") or ""
    out.append(f"<footer>{esc(footer)}{' · ' if footer and contact and contact not in footer else ''}{esc(contact if contact not in footer else '')}"
               f"<br>Generated {dt.date.today().isoformat()} from the recalculated workbook. Review before sharing: this report contains client names and activity.</footer>")
    out.append("</main></body></html>")
    return "".join(out)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", required=True)
    ap.add_argument("--workbook", default="")
    ap.add_argument("--out", default="")
    ap.add_argument("--theme", default="", choices=["", "auto", "light", "dark"])
    a = ap.parse_args()
    config_dir = os.path.dirname(os.path.abspath(a.config))
    cfg = json.load(open(a.config, encoding="utf-8"))

    def path(p: str) -> str:
        return p if os.path.isabs(p) else os.path.join(config_dir, p)

    workbook = path(a.workbook or cfg["output_xlsx"])
    out = path(a.out or os.path.splitext(cfg["output_xlsx"])[0] + "_report.html")
    document = build(cfg, workbook, a.theme, config_dir)
    with open(out, "w", encoding="utf-8") as handle:
        handle.write(document)
    print(f"saved {out} ({len(document)} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
