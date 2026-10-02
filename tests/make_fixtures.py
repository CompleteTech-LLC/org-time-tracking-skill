#!/usr/bin/env python3
"""Synthetic end-to-end fixture suite: tree -> workbook -> report, neutral and CompleteTech branding.

Uses a fictional organization and synthetic people only; writes under tests/out/ (git-ignored). Prints ALL OK on success.
The workbook's formulas are not evaluated here (that needs Excel or LibreOffice); they are checked structurally, and the
report renderer is exercised against a values-only stub with the same layout.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "out"
PY = sys.executable

LINES = """\
2026-07-01||staff|mail|Riley Chen|Weekly sync notes and open questions
2026-07-02||related|mail|Pat Example|Quote request for 5 licences
2026-07-02||alerts|CPU alert|noreply@northwind.example|CPU usage exceeds 90% threshold
2026-07-02||alerts|DB log alert|noreply@northwind.example|Transaction log file size crossed threshold
2026-07-01|09:05|teams|Riley Chen (1:1)|Riley|Can you check ticket 4411 before the call?
2026-07-01|09:08|teams|Riley Chen (1:1)|Pat|Looking now.
2026-07-01|09:20|teams|Riley Chen (1:1)|Riley|The admin login is Zq8$kd92mXp1 please use it
2026-07-06|10:00|teams|Riley Chen (1:1)|Riley|Planning at 1?
2026-07-06|10:02|teams|Riley Chen (1:1)|Pat|Yes.
2026-07-15|12:58|teams|Weekly sync|Riley|joining now
2026-07-15|13:01|teams|Weekly sync|Pat|same
2026-07-09|15:00|teams|Other chat|Sam|nothing from me here
"""
MEETINGS = [
    {"title": "Weekly sync", "group": "work", "start": "13:00", "end": "14:00", "dates": ["2026-07-01", "2026-07-08"]},
    {"title": "Weekly sync", "group": "work", "start": "13:00", "end": "14:00", "dates": ["2026-07-15"], "actual_min": 40,
     "evidence": "Teams: ended after 40m"},
    {"title": "Planning", "group": "work", "start": "13:00", "end": "13:30", "dates": ["2026-07-06"]},
    {"title": "Planning", "group": "work", "start": "13:00", "end": "13:30", "dates": ["2026-07-13"], "status": "Cancelled",
     "evidence": "cancellation email"},
    {"title": "Gym block", "group": "unconfirmed", "start": "17:00", "end": "18:00", "dates": ["2026-07-02", "2026-07-09"]},
]
TASKS = [
    {"task": "Fix tickets that re-open after closing", "source": "Teams", "requested_by": "Riley", "raised": "2026-07-01",
     "status": "Open", "low_h": 2, "high_h": 4, "basis": "several instances", "evidence": "Requested"},
    {"task": "Reduce alert noise", "source": "Email alerts", "requested_by": "-", "raised": "2026-07-02",
     "status": "Suggested, not requested", "low_h": 1, "high_h": 2, "basis": "mail rules", "evidence": "Suggested"},
]


def config(out_name: str, branding: dict, **billing) -> dict:
    return {
        "org": {"name": "Northwind Trading Co.", "domain": "northwind.example"},
        "tree_root": "./tree", "output_xlsx": f"./{out_name}.xlsx",
        "period": {"start": "2026-07-01", "end": "2026-07-31", "timezone": "America/New_York"},
        "me": {"teams_names": ["Pat"], "mail_markers": ["Pat Example"], "draft_prefix": "[Draft]", "calendar_subject_prefixes": ["Meeting"]},
        "alerts": {"sender_label": "noreply@northwind.example"},
        "billing": {"min_lapse_min": 15, "increment_min": 15, "email_actual_min": 5, "session_gap_min": 20, "session_padding_min": 2,
                    "meeting_day_min": 240, "meeting_basis": "Logged", "include_unconfirmed": "No", **billing},
        "meetings": {"work_group_label": "Northwind work", "unconfirmed_group_label": "Unconfirmed (personal?)", "held_rule": "teams_conversation_same_day"},
        "meetings_file": "meetings.json", "tasks_file": "tasks.json", "branding": branding,
        "style": {"font": "", "calendar_color_scale": {"mid_min": 60, "max_min": 240}, "minute_unit": "m", "theme": "auto"},
    }


def run(*args: str) -> str:
    result = subprocess.run([PY, *args], cwd=OUT, capture_output=True, text=True)
    if result.returncode:
        sys.stderr.write(result.stdout + result.stderr)
        raise SystemExit(f"command failed: {' '.join(args)}")
    return result.stdout


def check(condition: bool, message: str) -> None:
    if not condition:
        raise SystemExit("FAIL: " + message)
    print("ok:", message)


def stub_workbook(path: Path) -> None:
    """Values-only workbook with the layout build_report.py reads (no formulas, so cached values exist)."""
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s["A5"], s["B5"] = "Metric", "Value"
    for i, (k, v) in enumerate([("Email: staff", 1), ("Teams messages", 12), ("My time in communications (h), logged basis", 6.5)], 6):
        s.cell(row=i, column=1, value=k)
        s.cell(row=i, column=2, value=v)
    t = wb.create_sheet("Time Spent")
    for r, (k, c, m) in enumerate([("Teams chat sessions", 3, 60), ("Emails written", 1, 15), ("Held meetings", 2, 80), ("Top-up", 2, 400),
                                   ("Unconfirmed", 0, 0), ("Total", None, 555), ("Best", 4, 700), ("Best total", None, 900)], 2):
        t.cell(row=r, column=1, value=k)
        t.cell(row=r, column=2, value=c)
        t.cell(row=r, column=4, value=m)
    d = wb.create_sheet("Daily Time")
    import datetime as dt
    d.append(["Date", "Total (min)"])
    d.append([dt.datetime(2026, 7, 1), 255])
    d.append([dt.datetime(2026, 7, 6), 240])
    m_ = wb.create_sheet("Meetings")
    m_.append(["Date", "Meeting", "Group", "Status", "Teams actual (min)", "Evidence"])
    m_.append([dt.datetime(2026, 7, 1), "Weekly sync", "Northwind work", "Held (conversation that day)", None, "calendar"])
    m_.append([dt.datetime(2026, 7, 8), "Weekly sync", "Northwind work", "Not held (no conversation that day)", None, "rule"])
    k = wb.create_sheet("Tasks")
    k.append(["#", "Task", "Status", "Est. low (h)", "Est. high (h)", "Est. likely (h)"])
    k.append([1, "Fix tickets", "Open", 2, 4, 3])
    wb.save(path)


def main() -> int:
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True)
    (OUT / "lines.txt").write_text(LINES, encoding="utf-8")
    (OUT / "meetings.json").write_text(json.dumps(MEETINGS), encoding="utf-8")
    (OUT / "tasks.json").write_text(json.dumps(TASKS), encoding="utf-8")

    print(run(str(ROOT / "scripts/build_tree.py"), "--in", "lines.txt", "--out", "tree", "--start", "2026-07-01", "--tz", "America/New_York").strip())
    teams = (OUT / "tree/teams/2026-07/01.md").read_text(encoding="utf-8")
    check("Zq8" not in teams and "[REDACTED]" in teams, "credential-like token redacted in the tree")
    check(not any("Zq8" in p.read_text(encoding="utf-8") for p in (OUT / "tree").rglob("*.md")), "no secret fragment anywhere in the tree")

    neutral = config("neutral", {"preset": "neutral"})
    cthc = config("brand", {"preset": "completetech"}, min_lapse_min=30)
    for name, cfg in (("neutral", neutral), ("brand", cthc)):
        (OUT / f"{name}.json").write_text(json.dumps(cfg), encoding="utf-8")
        print(run(str(ROOT / "scripts/make_workbook.py"), "--config", f"{name}.json").strip())

    wn, wbk = load_workbook(OUT / "neutral.xlsx"), load_workbook(OUT / "brand.xlsx")
    expected = {"Summary", "Communications", "Teams Sessions", "Teams Messages", "Tasks", "Meetings", "Time Spent", "Daily Time", "Day Detail", "Calendar View", "Alert Rollup"}
    check(expected <= set(wn.sheetnames), "all workbook tabs are present")

    ts_n, ts_b = wn["Time Spent"], wbk["Time Spent"]
    check((ts_n["I3"].value, ts_n["I4"].value, ts_n["I7"].value) == (15, 15, 240), "billing parameters flow from the config into Time Spent inputs")
    check(ts_b["I3"].value == 30, "a different config value changes the input (min_lapse_min=30)")
    check(str(ts_n["I5"].value).startswith("=CEILING(MAX("), "billed minutes per email is a formula over the inputs")

    statuses = {(r[0].value.date().isoformat(), r[2].value): r[7].value for r in wn["Meetings"].iter_rows(min_row=2) if hasattr(r[0].value, "date")}
    check(statuses[("2026-07-01", "Weekly sync")].startswith("Held"), "meeting with Teams conversation that day is held")
    check(statuses[("2026-07-08", "Weekly sync")].startswith("Not held"), "meeting with no conversation that day is not held")
    check(statuses[("2026-07-15", "Weekly sync")].startswith("Held (Teams record"), "meeting with a card and conversation is held with the card")
    check(statuses[("2026-07-13", "Planning")] == "Cancelled", "explicit status wins over the rule")

    check(wn["Summary"]["A1"].value in (None, ""), "neutral workbook carries no company eyebrow")
    check(wn["Summary"].cell(row=5, column=1).fill.fgColor.rgb.endswith("1F3864"), "neutral accent on table headers")
    check(wbk["Summary"]["A1"].value == "COMPLETETECH LLC", "CompleteTech preset sets the eyebrow")
    check(wbk["Summary"].cell(row=5, column=1).fill.fgColor.rgb.endswith("1E3A8A"), "CompleteTech accent on table headers")
    check("CompleteTech LLC" in (wbk["Summary"].oddFooter.left.text or ""), "CompleteTech footer on printed pages")
    try:
        import PIL  # noqa: F401
        check(len(wbk["Summary"]._images) == 1 and not wn["Summary"]._images, "logo embedded only for the approved identity")
    except ImportError:
        print("skip: Pillow not installed, logo embedding not checked")

    stub_workbook(OUT / "stub.xlsx")
    (OUT / "report.json").write_text(json.dumps({**cthc, "output_xlsx": "./stub.xlsx"}), encoding="utf-8")
    run(str(ROOT / "scripts/build_report.py"), "--config", "report.json", "--out", "report_brand.html")
    (OUT / "report_n.json").write_text(json.dumps({**neutral, "output_xlsx": "./stub.xlsx"}), encoding="utf-8")
    run(str(ROOT / "scripts/build_report.py"), "--config", "report_n.json", "--out", "report_neutral.html", "--theme", "dark")
    page, neutral_page = (OUT / "report_brand.html").read_text(encoding="utf-8"), (OUT / "report_neutral.html").read_text(encoding="utf-8")
    check("Content-Security-Policy" in page and not re.search(r"""(?:src|href)\s*=\s*["']https?://""", page), "report is self-contained with a CSP")
    check("--accent:#1E3A8A" in page and "data:image/png;base64," in page and "COMPLETETECH LLC" in page, "branded report: accent, inlined logo, eyebrow")
    check("COMPLETETECH" not in neutral_page and "data:image" not in neutral_page and 'data-theme="dark"' in neutral_page,
          "neutral report: no company identity; theme override applied")
    check("prefers-color-scheme:dark" in page and "Held (conversation that day)" in page and "Not held" in page, "light/dark tokens and meeting table rendered")

    for bad in ({"preset": "no-such-brand"}, {"preset": "neutral", "accent": "blue"}, {"preset": "neutral", "logo": "https://example.com/x.png"}):
        (OUT / "bad.json").write_text(json.dumps(config("bad", bad)), encoding="utf-8")
        result = subprocess.run([PY, str(ROOT / "scripts/make_workbook.py"), "--config", "bad.json"], cwd=OUT, capture_output=True, text=True)
        check(result.returncode != 0, f"invalid branding is rejected: {bad}")

    print("ALL OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
